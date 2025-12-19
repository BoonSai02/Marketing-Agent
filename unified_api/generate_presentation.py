
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()

    # Define Brand Colors
    PRIMARY_COLOR = RGBColor(0, 51, 102)  # Dark Blue
    ACCENT_COLOR = RGBColor(255, 102, 0)  # Orange
    TEXT_COLOR = RGBColor(51, 51, 51)     # Dark Gray

    def set_title_format(title_shape):
        title_shape.text_frame.paragraphs[0].font.size = Pt(44)
        title_shape.text_frame.paragraphs[0].font.name = 'Arial'
        title_shape.text_frame.paragraphs[0].font.bold = True
        title_shape.text_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR

    def set_body_format(body_shape):
        tf = body_shape.text_frame
        for p in tf.paragraphs:
            p.font.size = Pt(18)
            p.font.name = 'Arial'
            p.font.color.rgb = TEXT_COLOR

    # ========== SLIDE 1: Title Slide ==========
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "AI Marketing Strategist"
    subtitle.text = "Your 24/7 Marketing Expert\nDelivering Data-Driven Growth Strategies in Minutes"
    
    # ========== SLIDE 2: The Problem ==========
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]

    title.text = "The Business Challenge"
    set_title_format(title)
    
    body.text = "Small and medium businesses struggle with:"
    
    tf = body.text_frame
    p = tf.add_paragraph()
    p.text = "Time Constraints: No bandwidth to research markets and competitors"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Budget Limitations: Marketing consultants cost $5,000-$50,000+ per project"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Uncertainty: Difficult to know which strategies will actually work"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Execution Gap: Great ideas without actionable implementation plans"
    p.level = 1
    
    set_body_format(body)

    # ========== SLIDE 3: Our Solution ==========
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]

    title.text = "What We Offer"
    set_title_format(title)

    body.text = "An intelligent business advisor that:"
    
    tf = body.text_frame
    p = tf.add_paragraph()
    p.text = "Researches your market and competitors in real-time"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Creates customized 90-day marketing roadmaps"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Provides actionable weekly steps with proven tactics"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Delivers instantly through a simple chat interface"
    p.level = 1

    set_body_format(body)

    # ========== SLIDE 4: Who Is This For? ==========
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]

    title.text = "Who Is This For?"
    set_title_format(title)

    body.text = "Perfect for:"
    
    tf = body.text_frame
    p = tf.add_paragraph()
    p.text = "Startup Founders: Need strategic marketing on a lean budget"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Small Business Owners: Want to grow without hiring expensive agencies"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Product Managers: Launching new products and need go-to-market plans"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Entrepreneurs: Building their brand from the ground up"
    p.level = 1

    set_body_format(body)

    # ========== SLIDE 5: Why Choose Us? ==========
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]
    
    title.text = "Why Choose Our Platform?"
    set_title_format(title)
    
    body.text = "Competitive advantages:"
    
    tf = body.text_frame
    p = tf.add_paragraph()
    p.text = "Cost Effective: 100x cheaper than hiring a marketing consultant"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Fast Results: Get your strategy in minutes, not weeks"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Data-Driven: Uses real market research, not generic templates"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Always Available: 24/7 access whenever inspiration strikes"
    p.level = 1
    
    set_body_format(body)

    # ========== SLIDE 6: How It Solves Your Problem ==========
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]

    title.text = "How It Works"
    set_title_format(title)

    body.text = "Simple 3-step process:"
    
    tf = body.text_frame
    p = tf.add_paragraph()
    p.text = "1. Tell us about your product or service through a conversation"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "2. Our system researches your industry, competitors, and opportunities"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "3. Receive a complete 90-day strategy with weekly action items"
    p.level = 1
    p = tf.add_paragraph()
    p.text = ""
    p.level = 1
    p = tf.add_paragraph()
    p.text = "No technical knowledge required. Just describe your business."
    p.level = 0
    p.font.bold = True

    set_body_format(body)

    # ========== SLIDE 7: Value Proposition ==========
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]
    
    title.text = "The Bottom Line"
    set_title_format(title)
    
    body.text = "What you get:"
    
    tf = body.text_frame
    p = tf.add_paragraph()
    p.text = "Save Time: Hours of research condensed into minutes"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Save Money: Professional-grade strategy at a fraction of the cost"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Reduce Risk: Make informed decisions backed by real data"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Grow Faster: Implement proven tactics immediately"
    p.level = 1
    
    set_body_format(body)

    # ========== SLIDE 8: Call to Action ==========
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "Ready to Grow Your Business?"
    subtitle.text = "Start your free strategy session today\n\nAI Marketing Strategist - Your Partner in Growth"

    # Save
    prs.save('AI_Marketing_Agent_Business_Pitch.pptx')
    print("✓ Presentation saved as 'AI_Marketing_Agent_Business_Pitch.pptx'")

if __name__ == "__main__":
    create_presentation()
